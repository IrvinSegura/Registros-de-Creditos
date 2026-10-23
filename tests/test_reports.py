import io

import pytest
from openpyxl import load_workbook

from app import clock

URL = "/api/v1/credits"


@pytest.fixture
def portfolio(client, create_credit):
    ana = create_credit(client_name="Ana Paz", amount=1000, interest_rate=10, term_months=12, issue_date="2026-01-10")
    create_credit(client_name="Bruno Díaz", amount=5000, interest_rate=20, term_months=24, issue_date="2026-02-20")
    client.post(f"{URL}/{ana['id']}/payments", json={"amount": "250", "paid_on": "2026-03-01", "note": "=1+1"})
    return ana


def test_excel_export_has_the_credits_and_a_summary(client, portfolio):
    response = client.get(f"{URL}/export?format=xlsx&sort=amount&order=desc")

    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["Content-Disposition"] == f'attachment; filename="creditos-{clock.today().isoformat()}.xlsx"'
    workbook = load_workbook(io.BytesIO(response.data))
    sheet = workbook["Créditos"]
    assert sheet["A4"].value == "ID"
    assert [sheet.cell(row, 2).value for row in (5, 6)] == ["Bruno Díaz", "Ana Paz"]
    assert sheet["C5"].number_format == '"$"#,##0.00'
    assert sheet["C7"].value == "=SUM(C5:C6)"
    assert "Resumen" in workbook.sheetnames


def test_excel_export_describes_the_filters(client, portfolio):
    response = client.get(f"{URL}/export?format=xlsx&client_name=ana&status=overdue")

    sheet = load_workbook(io.BytesIO(response.data))["Créditos"]
    assert "Cliente contiene: ana; Estado: Atrasado" in sheet["A2"].value


def test_pdf_export(client, portfolio):
    response = client.get(f"{URL}/export?format=pdf")

    assert response.mimetype == "application/pdf"
    assert response.data.startswith(b"%PDF")


def test_pdf_export_without_matches(client, portfolio):
    assert client.get(f"{URL}/export?format=pdf&client_name=nadie").data.startswith(b"%PDF")


def test_unknown_export_format_is_rejected(client):
    assert client.get(f"{URL}/export?format=docx").status_code == 422


def test_statement_as_pdf(client, portfolio):
    response = client.get(f"{URL}/{portfolio['id']}/statement")

    assert response.mimetype == "application/pdf"
    assert response.data.startswith(b"%PDF")
    assert f"estado-de-cuenta-{portfolio['id']}-" in response.headers["Content-Disposition"]


def test_statement_as_excel_keeps_notes_as_text(client, portfolio):
    response = client.get(f"{URL}/{portfolio['id']}/statement?format=xlsx")

    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.sheetnames == ["Estado de cuenta", "Amortización", "Pagos"]
    payments = workbook["Pagos"]
    assert payments["B2"].value == 250
    # A note that looks like a formula stays text
    assert payments["C2"].value == "=1+1"
    assert payments["C2"].data_type == "s"
    assert workbook["Amortización"]["G2"].value in {"Pagada", "Vencida", "Pendiente"}


def test_viewers_can_download_reports(client_as, portfolio):
    viewer = client_as("viewer")

    assert viewer.get(f"{URL}/export?format=xlsx").status_code == 200
    assert viewer.get(f"{URL}/{portfolio['id']}/statement").status_code == 200
