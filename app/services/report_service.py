"""Excel and PDF reports: the portfolio list and each credit's account statement. Knows nothing about HTTP."""

import io
from collections.abc import Iterable
from datetime import date
from decimal import Decimal
from typing import Any

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app import clock
from app.models import Credit
from app.services import amortization
from app.services.credit_service import EXPORT_COLUMNS, STATUS_LABELS, export_rows, status_of

# Everything in the reports is read by end users, so it stays in Spanish
INSTALLMENT_LABELS = {amortization.PAID: "Pagada", amortization.OVERDUE: "Vencida", amortization.PENDING: "Pendiente"}
FILTER_LABELS = {
    "client_name": "Cliente contiene",
    "status": "Estado",
    "issue_date_from": "Otorgado desde",
    "issue_date_to": "Otorgado hasta",
    "amount_min": "Monto mínimo",
    "amount_max": "Monto máximo",
}

# Excel number formats
MONEY_FORMAT = '"$"#,##0.00'
FORMATS = {"money": MONEY_FORMAT, "percent": '0.00" %"', "integer": "0", "date": "dd/mm/yyyy", "text": "@"}
HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
MUTED_FONT = Font(color="666666")

# PDF colors (RGB)
PDF_PRIMARY = (31, 58, 95)
PDF_MUTED = (102, 102, 102)
PDF_ROW_FILL = (242, 245, 249)


def portfolio_report(credits: list[Credit], file_format: str, filters: dict[str, Any]) -> bytes:
    if file_format == "xlsx":
        return _portfolio_xlsx(credits, filters)
    return _portfolio_pdf(credits, filters)


def statement(credit: Credit, file_format: str) -> bytes:
    data = _statement_data(credit)
    if file_format == "xlsx":
        return _statement_xlsx(data)
    return _statement_pdf(data)


# ---------- Shared data ----------

def _totals(credits: list[Credit]) -> list[tuple[str, str]]:
    standings = [credit.standing for credit in credits]
    by_status: dict[str, int] = {}
    for credit in credits:
        by_status[status_of(credit)] = by_status.get(status_of(credit), 0) + 1
    # "1 atrasado", "12 vigentes"
    status_text = ", ".join(
        f"{count} {STATUS_LABELS[status].lower()}{'' if count == 1 else 's'}" for status, count in sorted(by_status.items())
    )
    return [
        ("Créditos", f"{len(credits)}" + (f" ({status_text})" if status_text else "")),
        ("Monto otorgado", _money(sum((credit.amount for credit in credits), Decimal(0)))),
        ("Pagado", _money(sum((standing.paid_total for standing in standings), Decimal(0)))),
        ("Saldo pendiente", _money(sum((standing.outstanding_balance for standing in standings), Decimal(0)))),
        ("Monto vencido", _money(sum((standing.overdue_amount for standing in standings), Decimal(0)))),
    ]


def _describe_filters(filters: dict[str, Any]) -> str:
    parts = []
    for name, label in FILTER_LABELS.items():
        value = filters.get(name)
        if value in (None, ""):
            continue
        if name == "status":
            value = STATUS_LABELS[value]
        elif isinstance(value, date):
            value = _date(value)
        elif isinstance(value, Decimal):
            value = _money(value)
        parts.append(f"{label}: {value}")
    return "; ".join(parts) if parts else "Sin filtros"


def _statement_data(credit: Credit) -> dict[str, Any]:
    standing, today = credit.standing, clock.today()
    schedule = amortization.build_schedule(credit.amount, credit.interest_rate, credit.term_months, credit.issue_date)
    next_payment = "Sin pagos pendientes"
    if standing.next_due_date:
        next_payment = f"{_date(standing.next_due_date)} · {_money(standing.next_payment)}"
    return {
        "title": f"Estado de cuenta · {credit.client_name}",
        "subtitle": f"Crédito #{credit.id} · generado el {_date(today)}",
        "terms": [
            ("Monto", _money(credit.amount)),
            ("Tasa anual", f"{credit.interest_rate:.2f} %"),
            ("Plazo", f"{credit.term_months} meses"),
            ("Fecha de otorgamiento", _date(credit.issue_date)),
            ("Cuota mensual", _money(standing.monthly_payment)),
            ("Estado", STATUS_LABELS[status_of(credit)]),
        ],
        "standing": [
            ("Pagado", _money(standing.paid_total)),
            ("Por pagar (con intereses)", _money(standing.remaining_total)),
            ("Saldo de capital", _money(standing.outstanding_balance)),
            ("Cuotas pagadas", f"{standing.installments_paid} de {credit.term_months}"),
            ("Monto vencido", _money(standing.overdue_amount)),
            ("Días de atraso", str(standing.days_overdue)),
            ("Próximo pago", next_payment),
        ],
        "schedule": [
            (
                installment.number,
                installment.due_date,
                installment.payment,
                installment.interest,
                installment.principal,
                installment.balance,
                INSTALLMENT_LABELS[amortization.installment_status(installment, standing, today)],
            )
            for installment in schedule
        ],
        "payments": [
            (payment.paid_on, payment.amount, payment.note, payment.user.full_name if payment.user else "")
            for payment in sorted(credit.payments, key=lambda payment: (payment.paid_on, payment.id))
        ],
    }


SCHEDULE_HEADERS = ("#", "Fecha", "Cuota", "Interés", "Capital", "Saldo", "Estado")
SCHEDULE_KINDS = ("integer", "date", "money", "money", "money", "money", "text")
PAYMENT_HEADERS = ("Fecha", "Monto", "Nota", "Registrado por")
PAYMENT_KINDS = ("date", "money", "text", "text")


# ---------- Excel ----------

def _portfolio_xlsx(credits: list[Credit], filters: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Créditos"
    sheet["A1"] = "Cartera de créditos"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = f"Generado el {_date(clock.today())} · {_describe_filters(filters)}"
    sheet["A2"].font = MUTED_FONT

    headers = [header for header, _, _ in EXPORT_COLUMNS]
    kinds = [kind for _, _, kind in EXPORT_COLUMNS]
    first_row = 4
    last_row = _write_table(sheet, first_row, headers, kinds, export_rows(credits))

    # Totals under the money columns that add up (not the monthly payment)
    total_row = last_row + 1
    sheet.cell(total_row, 2, "Total").font = Font(bold=True)
    for index, header in enumerate(headers, start=1):
        if header in ("Monto", "Pagado", "Saldo pendiente", "Monto vencido") and credits:
            column = get_column_letter(index)
            cell = sheet.cell(total_row, index, f"=SUM({column}{first_row + 1}:{column}{last_row})")
            cell.number_format = MONEY_FORMAT
            cell.font = Font(bold=True)

    sheet.freeze_panes = sheet.cell(first_row + 1, 1)
    if credits:
        sheet.auto_filter.ref = f"A{first_row}:{get_column_letter(len(headers))}{last_row}"

    summary = workbook.create_sheet("Resumen")
    summary["A1"] = "Resumen de la cartera"
    summary["A1"].font = TITLE_FONT
    for offset, (label, value) in enumerate(_totals(credits), start=3):
        summary.cell(offset, 1, label).font = Font(bold=True)
        _set(summary.cell(offset, 2), value)
    _fit_columns(summary)
    return _save(workbook)


def _statement_xlsx(data: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Estado de cuenta"
    sheet["A1"] = data["title"]
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = data["subtitle"]
    sheet["A2"].font = MUTED_FONT

    row = 4
    for heading, pairs in (("Datos del crédito", data["terms"]), ("Situación al día de hoy", data["standing"])):
        sheet.cell(row, 1, heading).font = Font(bold=True, size=12)
        row += 1
        for label, value in pairs:
            sheet.cell(row, 1, label).font = Font(bold=True)
            _set(sheet.cell(row, 2), value)
            row += 1
        row += 1
    _fit_columns(sheet)

    schedule = workbook.create_sheet("Amortización")
    last = _write_table(schedule, 1, SCHEDULE_HEADERS, SCHEDULE_KINDS, data["schedule"])
    schedule.freeze_panes = "A2"
    schedule.auto_filter.ref = f"A1:{get_column_letter(len(SCHEDULE_HEADERS))}{last}"

    payments = workbook.create_sheet("Pagos")
    _write_table(payments, 1, PAYMENT_HEADERS, PAYMENT_KINDS, data["payments"])
    if not data["payments"]:
        payments.cell(2, 1, "Sin pagos registrados").font = MUTED_FONT
    payments.freeze_panes = "A2"
    return _save(workbook)


def _write_table(sheet: Worksheet, first_row: int, headers: Iterable[str], kinds: Iterable[str], rows: list) -> int:
    """Writes headers and rows starting at first_row; returns the last row written."""
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(first_row, column, header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row_number = first_row
    for row_number, row in enumerate(rows, start=first_row + 1):
        for column, (value, kind) in enumerate(zip(row, kinds, strict=True), start=1):
            cell = sheet.cell(row_number, column)
            _set(cell, value if value != "" else None)
            if kind != "text":
                cell.number_format = FORMATS[kind]
    _fit_columns(sheet, start_row=first_row)
    return row_number


def _set(cell, value: Any) -> None:
    cell.value = value
    # Text that starts with "=" would become a formula; notes are free text, so keep them as text
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _fit_columns(sheet: Worksheet, start_row: int = 1) -> None:
    widths: dict[int, int] = {}
    for row in sheet.iter_rows(min_row=start_row):
        for cell in row:
            if cell.value is None or (isinstance(cell.value, str) and cell.value.startswith("=SUM")):
                continue
            length = 12 if isinstance(cell.value, (date, Decimal, int, float)) else len(str(cell.value))
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 8), 50)


def _save(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------- PDF ----------

class _Report(FPDF):
    """A4 page with the app name on top and the page number at the bottom."""

    def __init__(self, orientation: str):
        super().__init__(orientation=orientation, format="A4")
        self.set_auto_page_break(auto=True, margin=15)
        self.set_margins(12, 12, 12)
        self.set_title("Registro de Créditos")
        self.set_creator("Registro de Créditos")

    def header(self) -> None:
        self.set_font("Helvetica", "B", 9)
        self.set_text_color(*PDF_PRIMARY)
        self.cell(0, 6, "Registro de Créditos", new_x="LMARGIN", new_y="NEXT")
        self.set_text_color(0, 0, 0)
        self.ln(2)

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(*PDF_MUTED)
        self.cell(0, 6, f"Página {self.page_no()} de {{nb}}", align="C")


PORTFOLIO_PDF_COLUMNS = (
    # (header, index in EXPORT_COLUMNS, width in mm, alignment)
    ("ID", 0, 10, "R"),
    ("Cliente", 1, 52, "L"),
    ("Monto", 2, 28, "R"),
    ("Tasa", 3, 15, "R"),
    ("Plazo", 4, 14, "R"),
    ("Otorgado", 5, 22, "C"),
    ("Estado", 6, 20, "C"),
    ("Pagado", 9, 28, "R"),
    ("Saldo", 10, 28, "R"),
    ("Vencido", 11, 26, "R"),
    ("Próx. venc.", 13, 22, "C"),
)


def _portfolio_pdf(credits: list[Credit], filters: dict[str, Any]) -> bytes:
    pdf = _Report("landscape")
    pdf.add_page()
    _pdf_title(pdf, "Cartera de créditos", f"Generado el {_date(clock.today())} · {_describe_filters(filters)}")
    _pdf_pairs(pdf, _totals(credits), columns=5)

    kinds = [kind for _, _, kind in EXPORT_COLUMNS]
    rows = [
        [_text(row[index], kinds[index]) for _, index, _, _ in PORTFOLIO_PDF_COLUMNS]
        for row in export_rows(credits)
    ]
    if rows:
        _pdf_table(pdf, [column[0] for column in PORTFOLIO_PDF_COLUMNS],
                   [column[2] for column in PORTFOLIO_PDF_COLUMNS], [column[3] for column in PORTFOLIO_PDF_COLUMNS], rows)
    else:
        _pdf_note(pdf, "Ningún crédito coincide con los filtros.")
    return bytes(pdf.output())


def _statement_pdf(data: dict[str, Any]) -> bytes:
    pdf = _Report("portrait")
    pdf.add_page()
    _pdf_title(pdf, data["title"], data["subtitle"])
    _pdf_heading(pdf, "Datos del crédito")
    _pdf_pairs(pdf, data["terms"], columns=3)
    _pdf_heading(pdf, "Situación al día de hoy")
    _pdf_pairs(pdf, data["standing"], columns=3)

    _pdf_heading(pdf, "Tabla de amortización")
    _pdf_table(
        pdf, SCHEDULE_HEADERS, (10, 26, 28, 28, 28, 32, 34), ("R", "C", "R", "R", "R", "R", "C"),
        [[_text(value, kind) for value, kind in zip(row, SCHEDULE_KINDS, strict=True)] for row in data["schedule"]],
    )

    _pdf_heading(pdf, "Pagos registrados")
    if data["payments"]:
        _pdf_table(
            pdf, PAYMENT_HEADERS, (26, 30, 86, 44), ("C", "R", "L", "L"),
            [[_text(value, kind) for value, kind in zip(row, PAYMENT_KINDS, strict=True)] for row in data["payments"]],
        )
    else:
        _pdf_note(pdf, "Sin pagos registrados.")
    return bytes(pdf.output())


def _pdf_title(pdf: FPDF, title: str, subtitle: str) -> None:
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 8, _latin1(title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(*PDF_MUTED)
    pdf.multi_cell(0, 5, _latin1(subtitle), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(3)


def _pdf_heading(pdf: FPDF, text: str) -> None:
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 11)
    pdf.set_text_color(*PDF_PRIMARY)
    pdf.cell(0, 7, _latin1(text), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)


def _pdf_pairs(pdf: FPDF, pairs: list[tuple[str, str]], columns: int) -> None:
    """Label above value, in a grid of `columns` per row."""
    width = pdf.epw / columns
    for start in range(0, len(pairs), columns):
        chunk = pairs[start:start + columns]
        pdf.set_font("Helvetica", "", 8)
        pdf.set_text_color(*PDF_MUTED)
        for label, _ in chunk:
            pdf.cell(width, 5, _latin1(label))
        pdf.ln(5)
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(0, 0, 0)
        for _, value in chunk:
            pdf.cell(width, 6, _latin1(value))
        pdf.ln(8)


def _pdf_table(pdf: FPDF, headers, widths, aligns, rows: list[list[str]]) -> None:
    pdf.set_font("Helvetica", "", 8)
    with pdf.table(
        col_widths=widths,
        width=sum(widths),
        align="LEFT",
        text_align=aligns,
        line_height=5,
        headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=PDF_PRIMARY),
        cell_fill_color=PDF_ROW_FILL,
        cell_fill_mode="ROWS",
        borders_layout="HORIZONTAL_LINES",
    ) as table:
        header_row = table.row()
        for header in headers:
            header_row.cell(_latin1(header))
        for row in rows:
            table_row = table.row()
            for value in row:
                table_row.cell(_latin1(value))


def _pdf_note(pdf: FPDF, text: str) -> None:
    pdf.set_font("Helvetica", "I", 9)
    pdf.set_text_color(*PDF_MUTED)
    pdf.cell(0, 6, _latin1(text), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)


# ---------- Formatting ----------

def _money(value: Decimal | None) -> str:
    return f"${value:,.2f}" if value is not None else ""


def _date(value: date | None) -> str:
    return value.strftime("%d/%m/%Y") if value else ""


def _text(value: Any, kind: str) -> str:
    if value in (None, ""):
        return ""
    if kind == "money":
        return _money(value)
    if kind == "percent":
        return f"{value:.2f} %"
    if kind == "date":
        return _date(value)
    return str(value)


def _latin1(text: str) -> str:
    # The built-in PDF fonts only cover Latin-1, which includes Spanish; anything else becomes "?"
    return text.encode("latin-1", errors="replace").decode("latin-1")
